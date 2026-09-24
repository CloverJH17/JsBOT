#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Pruebas nuevas de auditoría, conciliación y exportación multiformato."""
import csv
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from modulos import auditor_reportes as ar
from modulos.motor_export_auditoria import (
    clasificar_actividad_datos,
    parsear_csv_actividades_infoapp,
    parsear_csv_servicios_infoapp,
)


class FakePage:
    def __init__(self):
        self.html = ""

    def set_content(self, html, wait_until=None):
        self.html = html

    def pdf(self, path=None, **kwargs):
        Path(path).write_bytes(b"%PDF-1.4\n%%EOF\n")


class FakeBrowser:
    def __init__(self):
        self.page = FakePage()

    def new_page(self):
        return self.page

    def close(self):
        pass


class FakeChromium:
    def __init__(self):
        self.browser = FakeBrowser()

    def launch(self, **kwargs):
        return self.browser


class FakePlaywright:
    def __init__(self):
        self.chromium = FakeChromium()


class FakePlaywrightContext:
    def __init__(self):
        self.playwright = FakePlaywright()

    def __enter__(self):
        return self.playwright

    def __exit__(self, exc_type, exc, tb):
        return False


class TestAuditoriaReportesNueva(unittest.TestCase):
    def test_parser_csv_actividades_deduplica_y_clasifica(self):
        header = [
            "id", "info_id", "code_info", "user_id", "line_action", "report_type",
            "estate", "municipality", "parish", "activity_title", "date_ini",
            "person_fe", "person_ma", "responsible_name", "total_products", "tipo_taller",
        ]
        fila = ["A1", "I1", "S1", "U1", "Aprendizaje", "Formación", "Yaracuy", "San Felipe", "Centro", "Taller Python", "2026-09-01", "2", "3", "Facilitador", "0", "Robótica"]
        contenido = "\n".join(["|".join(header), "|".join(fila), "|".join(fila)]).encode("utf-8")
        resultado = parsear_csv_actividades_infoapp(contenido)
        self.assertEqual(len(resultado), 1)
        self.assertEqual(resultado[0]["participantes"], 5)
        self.assertEqual(resultado[0]["tipo_clasificacion"], "formacion")

    def test_parser_csv_servicios_acepta_esquema_minimo(self):
        header = ["id", "user_id", "info_id", "user_info_cod", "user_nombres", "user_apellidos", "user_dni", "user_genero", "user_edad", "user_profesion", "user_estado", "user_municipio", "user_tipo_servicio", "user_fecha_servicio"]
        fila = ["S1", "U1", "I1", "SRC1", "Ana", "Pérez", "25123456", "F", "30", "Docente", "Yaracuy", "San Felipe", "Capacitación", "2026-09-01"]
        contenido = "\n".join(["|".join(header), "|".join(fila)]).encode("utf-8")
        resultado = parsear_csv_servicios_infoapp(contenido)
        self.assertEqual(len(resultado), 1)
        self.assertEqual(resultado[0]["usuario"], "Ana Pérez")
        self.assertTrue(resultado[0]["cedulado"])

    def test_clasificacion_y_conciliacion_detectan_discrepancias(self):
        self.assertEqual(clasificar_actividad_datos("Producto", "", "Contenido", "", 0), "producto")
        resultado = {
            "total_actividades": 2,
            "total_estudiantes": 3,
            "total_servicios": 1,
            "formaciones": [{"id": "A1", "fecha": "2026-09-01", "participantes": 2}],
            "productos": [{"id": "A2", "fecha": "2026-09-01", "productos": 1}],
            "otras_actividades": [],
            "resumen_facilitadores": {"U1": {"total_act": 1, "estudiantes": 2, "servicios": 0}},
        }
        control = ar.conciliar_balance_auditoria(resultado, total_declarado_servidor=3)
        self.assertFalse(control["cuadra"])
        self.assertGreaterEqual(len(control["hallazgos"]), 1)

    def test_excel_genera_cuatro_hojas_y_autofiltro(self):
        resultado = {
            "criterio_tipo": "uid",
            "criterio_valor": "U1",
            "f_ini": "2026-09-01",
            "f_fin": "2026-09-30",
            "facilitador_principal": "Facilitador",
            "total_actividades": 1,
            "total_procesadas": 1,
            "total_estudiantes": 2,
            "total_servicios": 1,
            "formaciones": [{"id": "A1", "fecha": "2026-09-01", "uid": "U1", "info_id": "S1", "taller": "Python", "titulo": "Curso", "responsable": "Facilitador", "participantes": 2, "productos": 0}],
            "productos": [],
            "otras_actividades": [],
            "servicios": [{"fecha": "2026-09-02", "uid": "U1", "info_id": "S1", "servicio": "Capacitación", "cedula": "25123456", "id_usuario": "1", "usuario": "Ana", "profesion": "Docente"}],
            "resumen_facilitadores": {"U1": {"uid": "U1", "nombre": "Facilitador", "info_id": "S1", "formaciones": 1, "estudiantes": 2, "productos": 0, "otras": 0, "total_act": 1, "servicios": 1}},
        }
        with tempfile.TemporaryDirectory() as tmp:
            ruta = Path(tmp) / "reporte.xlsx"
            with patch.object(ar, "REPORTES_DIR", tmp):
                salida = ar.exportar_reporte_excel(resultado, str(ruta))
            self.assertEqual(Path(salida), ruta)
            workbook = load_workbook(salida, read_only=False)
            try:
                self.assertEqual(workbook.sheetnames, ["Resumen por Facilitador", "Actividades", "Servicios", "Resumen Ejecutivo"])
                self.assertEqual(workbook["Resumen por Facilitador"].auto_filter.ref, "A1:H5")
            finally:
                workbook.close()

    def test_pdf_escapa_contenido_html_in_controlado(self):
        resultado = {
            "criterio_tipo": "uid",
            "criterio_valor": "U1",
            "f_ini": "<script>alert('x')</script>",
            "f_fin": "2026-09-30",
            "facilitador_principal": "Facilitador",
            "total_actividades": 0,
            "total_estudiantes": 0,
            "total_servicios": 0,
            "formaciones": [],
            "productos": [],
            "otras_actividades": [],
            "servicios": [],
            "resumen_facilitadores": {},
        }
        with tempfile.TemporaryDirectory() as tmp:
            ruta = Path(tmp) / "reporte.pdf"
            fake = FakePlaywrightContext()
            with patch("playwright.sync_api.sync_playwright", return_value=fake):
                salida = ar.exportar_reporte_pdf(resultado, str(ruta))
            self.assertEqual(Path(salida), ruta)
            self.assertTrue(ruta.exists())
            captured_html = fake.playwright.chromium.browser.page.html
            self.assertNotIn("<script>alert", captured_html)
            self.assertIn("&lt;script&gt;", captured_html)


if __name__ == "__main__":
    unittest.main(verbosity=2)
