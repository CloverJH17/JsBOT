"""
Pruebas Unitarias para el Módulo Inspector de Auditoría (v4.2.0)
modulos/auditor_reportes.py
"""

import os
import unittest
from unittest.mock import patch, MagicMock
import openpyxl
from modulos import auditor_reportes as ar

class TestAuditorReportes(unittest.TestCase):

    def setUp(self):
        self.datos_prueba = {
            "exito": True,
            "criterio_tipo": "uid",
            "criterio_valor": "1325",
            "f_ini": "2026-08-01",
            "f_fin": "2026-08-31",
            "facilitador_principal": "Facilitador de Prueba",
            "total_actividades": 3,
            "total_procesadas": 3,
            "formaciones": [
                {
                    "fecha": "2026-08-10",
                    "id_activity": "101",
                    "titulo": "Taller de Robótica",
                    "dimensiones": "Comunidades de aprendizaje % Robótica Educativa % Tinkercad",
                    "area": "Robótica Educativa",
                    "taller": "Tinkercad",
                    "responsable": "Facilitador de Prueba",
                    "uid": "1325",
                    "participantes": 15,
                    "productos": 0,
                    "info_id": "NRYAR24"
                }
            ],
            "productos": [
                {
                    "fecha": "2026-08-15",
                    "id_activity": "102",
                    "titulo": "Video Tutorial Scratch",
                    "dimensiones": "Medios digitales % Producción de contenido",
                    "area": "Medios digitales",
                    "taller": "Producción de contenido",
                    "responsable": "Facilitador de Prueba",
                    "uid": "1325",
                    "participantes": 0,
                    "productos": 1,
                    "info_id": "NRYAR24"
                }
            ],
            "otras_actividades": [
                {
                    "fecha": "2026-08-20",
                    "id_activity": "103",
                    "titulo": "Soporte Comunitario",
                    "dimensiones": "Acompañamiento a comunidad",
                    "area": "Comunidad",
                    "taller": "Acompañamiento",
                    "responsable": "Facilitador de Prueba",
                    "uid": "1325",
                    "participantes": 0,
                    "productos": 0,
                    "info_id": "NRYAR24"
                }
            ],
            "total_estudiantes": 15,
            "total_servicios": 1,
            "servicios": [
                {
                    "fecha": "2026-08-12",
                    "servicio": "Navegación Asistida",
                    "cedula": "24123456",
                    "id_usuario": "501",
                    "usuario": "María Pérez",
                    "profesion": "Estudiante",
                    "uid": "1325",
                    "info_id": "NRYAR24"
                }
            ],
            "conteo_servicios": {"Navegación Asistida": 1},
            "cedulados_serv": 1,
            "no_cedulados_serv": 0,
            "resumen_facilitadores": {
                "1325": {
                    "nombre": "Facilitador de Prueba",
                    "info_id": "NRYAR24",
                    "formaciones": 1,
                    "estudiantes": 15,
                    "productos": 1,
                    "otras": 1,
                    "total_act": 3,
                    "servicios": 1
                }
            },
            "cuadre_perfecto": True
        }
        self.archivos_creados = []

    def tearDown(self):
        for arch in self.archivos_creados:
            if os.path.exists(arch):
                try:
                    os.remove(arch)
                except Exception:
                    pass

    def test_cargar_credenciales_auditoria_devuelve_tupla(self):
        usr, pwd = ar.cargar_credenciales_auditoria(rol_auditor=False)
        self.assertIsInstance(usr, str)
        self.assertIsInstance(pwd, str)

    def test_exportar_reporte_csv(self):
        ruta_csv = ar.exportar_reporte_auditoria(self.datos_prueba, formato="csv")
        self.assertTrue(os.path.exists(ruta_csv), f"El archivo CSV no fue creado en {ruta_csv}")
        self.archivos_creados.append(ruta_csv)
        self.assertTrue(ruta_csv.endswith(".csv"))

        with open(ruta_csv, mode="r", encoding="utf-8-sig") as f:
            contenido = f.read()
            self.assertIn("Titulo", contenido)
            self.assertIn("Facilitador de Prueba", contenido)
            self.assertIn("Taller de Robótica", contenido)

    def test_exportar_reporte_excel(self):
        ruta_xlsx = ar.exportar_reporte_auditoria(self.datos_prueba, formato="excel")
        self.assertTrue(os.path.exists(ruta_xlsx), f"El archivo Excel no fue creado en {ruta_xlsx}")
        self.archivos_creados.append(ruta_xlsx)
        self.assertTrue(ruta_xlsx.endswith(".xlsx"))

        # Validar estructura del libro openpyxl (v4.2.3)
        wb = openpyxl.load_workbook(ruta_xlsx)
        nombres_hojas = wb.sheetnames
        self.assertIn("Resumen por Facilitador", nombres_hojas)
        self.assertIn("Actividades", nombres_hojas)
        self.assertIn("Servicios", nombres_hojas)
        self.assertIn("Resumen Ejecutivo", nombres_hojas)

        # Hoja 1: Resumen por Facilitador debe ser la primera hoja y contener la sede
        ws_fac = wb["Resumen por Facilitador"]
        texto_sede_encontrado = False
        subtotal_encontrado = False
        for row in ws_fac.iter_rows(values_only=True):
            row_str = str(row)
            if "INFOCENTRO: NRYAR24" in row_str:
                texto_sede_encontrado = True
            if "Subtotal NRYAR24" in row_str:
                subtotal_encontrado = True
        self.assertTrue(texto_sede_encontrado, "No se encontró la cabecera de grupo de Infocentro")
        self.assertTrue(subtotal_encontrado, "No se encontró la fila de subtotal por Infocentro")

        # Validar contenido en Hoja Resumen Ejecutivo
        ws_resumen = wb["Resumen Ejecutivo"]
        texto_titulo = ws_resumen["A1"].value
        self.assertIn("AUDITORÍA", texto_titulo)

        # Validar hoja Actividades
        ws_act = wb["Actividades"]
        encontrado_taller = False
        for row in ws_act.iter_rows(values_only=True):
            if "Taller de Robótica" in str(row):
                encontrado_taller = True
                break
        self.assertTrue(encontrado_taller)
        wb.close()

    def test_balance_matematico_cuadre_perfecto(self):
        tot_act = self.datos_prueba["total_actividades"]
        tot_proc = len(self.datos_prueba["formaciones"]) + len(self.datos_prueba["productos"]) + len(self.datos_prueba["otras_actividades"])
        self.assertEqual(tot_act, tot_proc)
        self.assertTrue(self.datos_prueba["cuadre_perfecto"])

    @patch("modulos.auditor_reportes.iniciar_driver_auditoria")
    @patch("modulos.auditor_reportes.autenticar_infoapp")
    @patch("modulos.auditor_reportes.consultar_actividades_infoapp")
    @patch("modulos.auditor_reportes.consultar_servicios_infoapp")
    def test_firma_polimorfica_ejecutar_auditoria(self, mock_serv, mock_act, mock_auth, mock_driver):
        """Verifica que ejecutar_auditoria acepte kwargs (uid, start_at, finish_at, etc.) sin lanzar TypeError."""
        mock_act.return_value = (0, [])
        mock_serv.return_value = (0, [])
        mock_driver_instance = MagicMock()
        mock_driver.return_value = mock_driver_instance

        # Invocación con argumentos por palabra clave que antes fallaban
        resultado = ar.ejecutar_auditoria(
            uid="1325",
            info_id="NRYAR24",
            estado="Yaracuy",
            start_at="2026-08-01",
            finish_at="2026-08-31",
            rol_auditor=False,
            formato="excel"
        )
        self.assertTrue(resultado.get("exito"))
        self.assertEqual(resultado.get("criterio_tipo"), "uid")
        self.assertEqual(resultado.get("criterio_valor"), "1325")
        self.assertEqual(resultado.get("f_ini"), "2026-08-01")
        self.assertEqual(resultado.get("f_fin"), "2026-08-31")
        if resultado.get("archivo_exportado"):
            self.archivos_creados.append(resultado.get("archivo_exportado"))

    def test_parseo_bs4_actividades_y_servicios(self):
        """Verifica que el parser BeautifulSoup extraiga correctamente actividades y servicios."""
        html_act = """
        <table class="table table-hover">
            <tbody>
                <tr>
                    <td><span>2026-08-10</span></td>
                    <td><p class="id_activity" id="999">999</p><p class="data_titulo">Taller Python</p><p class="data_dimensiones">Comunidades de aprendizaje % Robótica % Python Básico</p><p class="responsible_name">Facilitador Uno</p><p class="user_id">1325</p></td>
                    <td>NRYAR24</td>
                    <td><a class="btn-info btn-sm">25</a></td>
                    <td><a class="btn-danger btn-sm">0</a></td>
                </tr>
            </tbody>
        </table>
        """
        acts = ar.parsear_pagina_actividades_bs4(html_act, default_info_id="NRYAR24", default_uid="1325")
        self.assertEqual(len(acts), 1)
        self.assertEqual(acts[0]["titulo"], "Taller Python")
        self.assertEqual(acts[0]["participantes"], 25)
        self.assertEqual(acts[0]["taller"], "Python Básico")
        self.assertEqual(acts[0]["area"], "Robótica")

        html_srv = """
        <table class="table table-bordered">
            <tbody>
                <tr>
                    <td>1325</td>
                    <td>2026-08-11</td>
                    <td>Facilitador</td>
                    <td>NRYAR24</td>
                    <td>Yaracuy</td>
                    <td>Asesoría Virtual</td>
                    <td>V-12345678</td>
                    <td>888</td>
                    <td>Juan Gómez</td>
                    <td>Docente</td>
                </tr>
            </tbody>
        </table>
        """
        srvs = ar.parsear_pagina_servicios_bs4(html_srv, default_info_id="NRYAR24", default_uid="1325")
        self.assertEqual(len(srvs), 1)
        self.assertEqual(srvs[0]["servicio"], "Asesoría Virtual")
        self.assertEqual(srvs[0]["cedula"], "V-12345678")
        self.assertEqual(srvs[0]["usuario"], "Juan Gómez")
        self.assertEqual(srvs[0]["profesion"], "Docente")

if __name__ == "__main__":
    unittest.main()
