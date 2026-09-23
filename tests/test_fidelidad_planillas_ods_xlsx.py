#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
===============================================================================
TEST DE FIDELIDAD INSTITUCIONAL Y COMPATIBILIDAD ODS / XLSX
(test_fidelidad_planillas_ods_xlsx.py)
===============================================================================
Sistema   : JsBOT (Robotic Process Automation)
Objetivo  : Validar que:
            1. Los archivos ODS generados respeten estrictamente el límite de filas
               de hojas de cálculo (<= 1.048.576) para no corromperse en Excel.
            2. El XML del archivo ODS conserve los namespaces canónicos ODF sin
               prefijos no estándar (ns1:, ns2:, etc.).
            3. Los archivos XLSX repliquen fielmente el formato institucional de
               la plantilla base (celdas combinadas B:C para nombres, membretes,
               metadatos, bordes y nota legal enmarcada).
            4. El despachador multiformato enrute correctamente según extensión.
===============================================================================
"""

import os
import sys
import re
import zipfile
import tempfile
import unittest
import openpyxl
import xml.etree.ElementTree as ET

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

import modulos.generador_planilla as gp


def _crear_participantes_muestra(cantidad: int = 57) -> list:
    """Genera una lista de participantes heterogéneos para pruebas."""
    lista = []
    for i in range(1, cantidad + 1):
        if i % 3 == 0:
            doc = {"cedulado": "escolar", "cedula_escolar": f"20120101{i:03d}"}
        elif i % 5 == 0:
            doc = {"cedulado": "no", "cedula_padre": f"18765{i:03d}"}
        else:
            doc = {"cedulado": "si", "cedula": f"{25000000 + i}"}

        p = {
            "nombre": f"Nombre{i}",
            "apellido": f"Apellido{i}",
            "nacimiento": "2000-01-15",
            "genero": "M" if i % 2 == 0 else "F",
            "direccion": f"Comunidad Sector {i}, San Felipe",
            "correo": f"participante{i}@ejemplo.com",
            "telefono": f"0412-{1000000 + i}",
            "edad": 15 + (i % 20),
            "nivel": "Educación Media General",
            "ocupacion": "Estudiante"
        }
        p.update(doc)
        lista.append(p)
    return lista


class TestFidelidadPlanillasODSyXLSX(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.participantes = _crear_participantes_muestra(57)
        self.url_prueba = "https://infoapp2.infocentro.gob.ve/admin/index.php?view=participants_list&id_activity=9999&activity=Taller%20de%20Python&estate=Yaracuy&code_info=Yar23&line_action=Tecnologias%20Libres&date_activity=2026-09-20/2026-09-25"

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_ods_no_excede_limite_filas_excel(self):
        """Valida que el archivo ODS generado no sobrepase el límite máximo de 1.048.576 filas."""
        ruta_ods = os.path.join(self.temp_dir.name, "prueba_filas.ods")
        gp.generar_planilla_ods_odfdo(self.participantes, id_actividad="9999", url_actividad=self.url_prueba, ruta_salida=ruta_ods)

        self.assertTrue(os.path.exists(ruta_ods), "El archivo ODS no fue creado.")

        with zipfile.ZipFile(ruta_ods, 'r') as z:
            xml_content = z.read('content.xml')

        root = ET.fromstring(xml_content)
        table_ns = 'urn:oasis:names:tc:opendocument:xmlns:table:1.0'
        rows = root.findall(f".//{{{table_ns}}}table-row")

        total_filas = 0
        for r in rows:
            rep = r.attrib.get(f"{{{table_ns}}}number-rows-repeated")
            if rep:
                total_filas += int(rep)
            else:
                total_filas += 1

        self.assertLessEqual(
            total_filas,
            1048576,
            f"El total de filas ({total_filas}) sobrepasa el límite absoluto de 1.048.576 filas permitido por Excel."
        )
        self.assertGreaterEqual(total_filas, 8 + 57 + 1, "Debe contener al menos cabeceras + 57 participantes + nota pie.")

    def test_ods_namespaces_canonicos_sin_prefijos_sintetizados(self):
        """Valida que content.xml no contenga prefijos sintéticos ns1: o ns2: que rompen Excel."""
        ruta_ods = os.path.join(self.temp_dir.name, "prueba_namespaces.ods")
        gp.generar_planilla_ods_odfdo(self.participantes, id_actividad="9999", url_actividad=self.url_prueba, ruta_salida=ruta_ods)

        with zipfile.ZipFile(ruta_ods, 'r') as z:
            xml_text = z.read('content.xml').decode('utf-8', errors='ignore')

        self.assertNotIn("ns1:", xml_text, "content.xml contiene el prefijo sintetizado ns1: incompatible con Excel.")
        self.assertNotIn("ns2:", xml_text, "content.xml contiene el prefijo sintetizado ns2: incompatible con Excel.")
        self.assertNotIn("ns0:", xml_text, "content.xml contiene el prefijo sintetizado ns0: incompatible con Excel.")
        self.assertIn("xmlns:table=", xml_text, "Falta el namespace canonical xmlns:table")
        self.assertIn("xmlns:office=", xml_text, "Falta el namespace canonical xmlns:office")
        self.assertIn("xmlns:style=", xml_text, "Falta el namespace canonical xmlns:style")

    def test_ods_preserva_nota_legal_pie(self):
        """Verifica que la nota legal institucional permanezca al pie de la tabla ODS."""
        ruta_ods = os.path.join(self.temp_dir.name, "prueba_nota.ods")
        gp.generar_planilla_ods_odfdo(self.participantes, id_actividad="9999", url_actividad=self.url_prueba, ruta_salida=ruta_ods)

        with zipfile.ZipFile(ruta_ods, 'r') as z:
            xml_text = z.read('content.xml').decode('utf-8', errors='ignore')

        self.assertIn("Nota: Para ser llenado con letra Imprenta", xml_text)

    def test_xlsx_fidelidad_institucional_y_celdas_combinadas(self):
        """Valida que el archivo XLSX contenga la estructura institucional 1:1 requerida."""
        ruta_xlsx = os.path.join(self.temp_dir.name, "prueba_fidelidad.xlsx")
        gp.generar_planilla_xlsx(self.participantes, id_actividad="9999", url_actividad=self.url_prueba, ruta_salida=ruta_xlsx)

        self.assertTrue(os.path.exists(ruta_xlsx), "El archivo XLSX no fue creado.")

        wb = openpyxl.load_workbook(ruta_xlsx)
        ws = wb.active

        # 1. Título y secciones
        self.assertEqual(ws.title, "Planilla de Inscripción")
        self.assertIn("Planilla de Inscripción", ws["A2"].value)
        self.assertIn("PLANILLA OFICIAL", ws["A2"].value)
        self.assertEqual(ws["A3"].value, "INFORMACIÓN GENERAL DEL PROCESO FORMATIVO")
        self.assertEqual(ws["A7"].value, "DATOS DE LOS PARTICIPANTES")

        # 2. Metadatos
        self.assertIn("Yaracuy", str(ws["A4"].value))
        self.assertIn("Yar23", str(ws["G4"].value))

        # 3. Cabecera tabla en fila 8
        self.assertEqual(ws["A8"].value, "N.º")
        self.assertEqual(ws["B8"].value, "Nombres y Apellidos*")
        self.assertEqual(ws["D8"].value, "Cedula *")
        self.assertEqual(ws["L8"].value, "Firma*")

        # 4. Validar celdas combinadas de nombres B:C para todos los participantes
        merged_ranges = [str(m) for m in ws.merged_cells.ranges]
        self.assertIn("B8:C8", merged_ranges, "La cabecera B8:C8 debe estar combinada.")

        # Verificar filas de participantes
        for r in range(9, 9 + len(self.participantes)):
            self.assertIn(f"B{r}:C{r}", merged_ranges, f"La fila {r} debe tener combinadas las columnas B y C para el nombre.")

        # 5. Validar datos de participantes
        self.assertEqual(ws["A9"].value, 1)
        self.assertIn("Nombre1", ws["B9"].value)
        self.assertIn("Apellido1", ws["B9"].value)

        # 6. Validar fila de nota legal al pie
        fila_pie = 9 + len(self.participantes)
        self.assertIn(f"A{fila_pie}:L{fila_pie}", merged_ranges, f"La nota al pie en fila {fila_pie} debe combinar A:L.")
        self.assertIn("Nota: Para ser llenado con letra Imprenta", ws[f"A{fila_pie}"].value)

    def test_despachador_multiformato(self):
        """Valida que generar_planilla_multiformato enrute correctamente según extensión."""
        ruta_base = os.path.join(self.temp_dir.name, "prueba_despacho")

        # Despacho por extensión XLSX
        salida_xlsx = gp.generar_planilla_multiformato(self.participantes, ruta_salida=ruta_base + ".xlsx")
        self.assertTrue(salida_xlsx.endswith(".xlsx"))
        self.assertTrue(os.path.exists(salida_xlsx))

        # Despacho por extensión ODS
        salida_ods = gp.generar_planilla_multiformato(self.participantes, ruta_salida=ruta_base + ".ods")
        self.assertTrue(salida_ods.endswith(".ods"))
        self.assertTrue(os.path.exists(salida_ods))


if __name__ == '__main__':
    unittest.main()
